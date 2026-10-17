# -*- coding: utf-8 -*-
import os
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from config import REPORT_DIR

def save_report(title, headers, rows):
    wb = Workbook()
    ws = wb.active
    ws.title = title
    ws.append(headers)

    # 表头样式
    for cell in ws[1]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='2C5F8A')
        cell.alignment = Alignment(horizontal='center', vertical='center')

    for row in rows:
        ws.append(list(row))

    # 状态为"异常"的行整行标红
    if '状态' in headers:
        col = headers.index('状态') + 1
        for r in ws.iter_rows(min_row=2):
            for c in r:
                c.alignment = Alignment(vertical='center', wrap_text=True)
            if r[col - 1].value == '异常':
                for c in r:
                    c.fill = PatternFill('solid', fgColor='FCE4E4')

    # 列宽
    widths = [18, 16, 14, 14, 50, 22]
    for i, w in enumerate(widths[:len(headers)], start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    fname = f'{title}_{datetime.now().strftime("%Y%m%d_%H%M%S")}.xlsx'
    path = os.path.join(REPORT_DIR, fname)
    wb.save(path)
    return path
