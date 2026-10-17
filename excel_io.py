# -*- coding: utf-8 -*-
from openpyxl import Workbook, load_workbook
from config import DEVICE_FILE

HEADERS = ['设备名称', '管理IP', '用户名', '密码', '设备类型']

def create_template():
    """生成设备清单模板 devices.xlsx"""
    wb = Workbook()
    ws = wb.active
    ws.title = '设备清单'
    ws.append(HEADERS)
    samples = [
        ['核心交换机SW1', '192.168.1.1', 'admin', 'Admin@123', '华为S5700'],
        ['汇聚交换机SW2', '192.168.1.2', 'admin', 'Admin@123', '华为S5700'],
        ['出口路由器R1', '192.168.1.254', 'admin', 'Admin@123', '华为AR2220'],
    ]
    for row in samples:
        ws.append(row)
    widths = [18, 16, 12, 16, 16]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[chr(64 + i)].width = w
    wb.save(DEVICE_FILE)
    print('已生成设备清单模板：', DEVICE_FILE)

def load_devices():
    """读取设备清单，返回字典列表"""
    wb = load_workbook(DEVICE_FILE)
    ws = wb.active
    devices = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row[0] or not row[1]:
            continue
        devices.append({
            'name': str(row[0]).strip(),
            'ip': str(row[1]).strip(),
            'username': str(row[2]).strip() if row[2] else 'admin',
            'password': str(row[3]).strip() if row[3] else '',
            'type': str(row[4]).strip() if row[4] else '',
        })
    return devices
